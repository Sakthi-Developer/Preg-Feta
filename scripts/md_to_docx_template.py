"""Render a manuscript Markdown file into the ProComm full-paper Word template.

Companion to `md_to_docx.py`, which produces a plain review copy in its own styling.
This one instead opens `Full-Paper-template.docx`, strips its instructional body, and
rebuilds the document out of the template's **own named styles** -- `paper title`,
`Author Data`, `Abstract`, `Heading 1/2`, `Body Text Indent`, `Caption`, `BulletList`,
`NumberedList`, `BlockQuote`, `Table Head`, `Table text` -- so the result inherits the
template's page size, margins, title-page footer and two-column section break rather
than imitating them.

    .venv/Scripts/python.exe scripts/md_to_docx_template.py CAFE_Hybrid_Manuscript.md

`[TABLE n HERE]` / `[FIGURE n HERE]` callouts are **left as placeholders**, marker and
caption intact, for the author to fill in by hand. Nothing is embedded and no artwork is
resolved from disk, so the artifacts stay under the author's control.

Heading numbers from the Markdown are kept verbatim. The template's own headings are
unnumbered, but this manuscript cross-references its sections by number ("see 4.4.1"),
so stripping them to match the template would break every one of those references.

Deliberately not a general Markdown implementation: anything it does not recognise is
emitted as plain text rather than silently dropped, so nothing disappears from a document
someone is about to submit.
"""
from __future__ import annotations

import argparse
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Twips
from docx.text.paragraph import Paragraph

REPO = Path(__file__).resolve().parents[1]

TEMPLATE = REPO / "Full-Paper-template.docx"

BODY_PT = 10        # template 'Normal'
TABLE_PT = 8        # tables are set smaller than body text to survive the column width
MONO = "Consolas"

# Printable width of the template's page, and of one of its two columns, in twips.
PAGE_TWIPS = 12240 - 2 * 1210
COL_TWIPS = (PAGE_TWIPS - 461) // 2

SUPERS = "¹²³⁴⁵⁶⁷⁸⁹"

# **bold**, *italic*, `code`, [text](url) -- captured so the split keeps the delimiters.
INLINE = re.compile(r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+`|\[[^\]]+\]\([^)]+\))")

# Backslash-escaped punctuation is parked on a sentinel before the inline pass, so an
# escaped asterisk cannot be mistaken for emphasis, then restored as a literal character.
_ESC = re.compile(r"\\([*_`|\[\]\\])")
_SENT = "\x00"


def _park(text: str) -> str:
    return _ESC.sub(lambda m: _SENT + str(ord(m.group(1))) + _SENT, text)


def _unpark(text: str) -> str:
    return re.sub(_SENT + r"(\d+)" + _SENT, lambda m: chr(int(m.group(1))), text)


def add_runs(par, text: str, size: float | None = None) -> None:
    """Emit `text` into `par`, honouring inline bold / italic / code / link markup.

    `size` overrides the paragraph style's font size -- needed for table cells and
    captions, which sit below the 10-point body text.
    """
    for tok in INLINE.split(_park(text)):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            run = par.add_run(_unpark(tok[2:-2]))
            run.bold = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            run = par.add_run(_unpark(tok[1:-1]))
            run.font.name = MONO
            run.font.size = Pt((size or BODY_PT) - 1)
            continue
        elif tok.startswith("[") and "](" in tok:
            run = par.add_run(_unpark(tok[1: tok.index("](")]))
            run.underline = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            run = par.add_run(_unpark(tok[1:-1]))
            run.italic = True
        else:
            run = par.add_run(_unpark(tok))
        if size is not None:
            run.font.size = Pt(size)


def split_row(line: str) -> list[str]:
    """Split a Markdown table row on unescaped pipes."""
    parts, cur, esc = [], "", False
    for ch in line.strip().strip("|"):
        if esc:
            cur += ch
            esc = False
        elif ch == "\\":
            esc = True
        elif ch == "|":
            parts.append(cur)
            cur = ""
        else:
            cur += ch
    parts.append(cur)
    return [p.strip() for p in parts]


def _sub(parent, tag: str, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn("w:" + k), str(v))
    parent.append(el)
    return el


# --------------------------------------------------------------------------- content

def para(doc, style: str, text: str = "", size: float | None = None,
         align=None, first_line: float | None = None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if first_line is not None:
        p.paragraph_format.first_line_indent = Inches(first_line)
    if text:
        add_runs(p, text, size)
    return p


def restart_numbering(doc, par) -> int:
    """Give `par` a fresh numbering sequence and return its numId.

    The template's NumberedList style points at one shared list, so every numbered list in
    the document continues the previous one -- section 3.1 came out as 7) 8) 9) 10). A
    restart needs a new `w:num` of its own, overriding the start value back to 1.
    """
    numbering = doc.part.numbering_part.element
    nums = numbering.findall(qn("w:num"))
    style_num = doc.styles["NumberedList"].element.find(
        qn("w:pPr")).find(qn("w:numPr")).find(qn("w:numId")).get(qn("w:val"))
    abstract = next(n.find(qn("w:abstractNumId")).get(qn("w:val"))
                    for n in nums if n.get(qn("w:numId")) == style_num)

    new_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    _sub(num, "w:abstractNumId", val=abstract)
    override = _sub(num, "w:lvlOverride", ilvl=0)
    _sub(override, "w:startOverride", val=1)
    numbering.append(num)                             # w:num elements follow w:abstractNum

    num_pr = par._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = new_id
    return new_id


def apply_numbering(par, num_id: int) -> None:
    """Attach a continuing item to the list started by `restart_numbering`."""
    num_pr = par._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


def add_placeholder(doc, text: str) -> None:
    """A `[TABLE n HERE]` / `[FIGURE n HERE]` callout, left for the author to fill in."""
    p = para(doc, "Caption", text, size=9)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)


def column_widths(header: list[str], body: list[list[str]], total: int) -> list[int]:
    """Distribute `total` twips across columns in proportion to their widest cell.

    Each column also gets a floor wide enough for its longest unbreakable word, so a short
    header like "Term" is not squeezed until "AUROC" wraps mid-acronym.
    """
    ncol = len(header)
    columns = [[re.sub(r"[*`\\]", "", (r[c] if c < len(r) else ""))
                for r in [header] + body] for c in range(ncol)]

    weights = [max(max(len(x) for x in col), 4) ** 0.75 for col in columns]
    #   ~115 twips per character of upper-case 8-point Times, plus the cell's side
    #   margins. Hyphens count as break opportunities, since Word wraps there.
    floors = [max(len(w) for w in re.split(r"[\s\-–—]+", " ".join(col)) or [""]) * 115 + 200
              for col in columns]
    if sum(floors) > total:                           # nothing fits: fall back to shares
        floors = [int(f * total / sum(floors)) for f in floors]

    scale = total / sum(weights)
    widths = [max(w * scale, f) for w, f in zip(weights, floors)]
    over = sum(widths) - total
    if over > 0:                                      # reclaim only from columns with slack
        slack = [w - f for w, f in zip(widths, floors)]
        pool = sum(slack)
        if pool > 0:
            widths = [w - min(over, pool) * s / pool for w, s in zip(widths, slack)]

    out = [int(round(w)) for w in widths]
    out[-1] += total - sum(out)                       # absorb rounding drift
    return out


def add_table(doc, rows: list[str], total_twips: int) -> None:
    """Render a Markdown pipe table using the template's table conventions."""
    header = split_row(rows[0])
    body = [split_row(r) for r in rows[2:]]           # rows[1] is the |---|---| rule
    ncol = len(header)
    widths = column_widths(header, body, total_twips)

    table = doc.add_table(rows=1, cols=ncol)

    # tblPr is rebuilt rather than patched: its children are schema-ordered, and appending
    # to whatever python-docx left behind produces a file Word offers to repair.
    tbl_pr = table._tbl.tblPr
    for child in list(tbl_pr):
        tbl_pr.remove(child)
    _sub(tbl_pr, "w:tblW", w=total_twips, type="dxa")
    _sub(tbl_pr, "w:jc", val="center")
    borders = _sub(tbl_pr, "w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        _sub(borders, "w:" + edge, val="single", sz=2, space=0, color="auto")
    _sub(tbl_pr, "w:tblLayout", type="fixed")
    margins = _sub(tbl_pr, "w:tblCellMar")
    for side, w in (("top", 20), ("left", 72), ("bottom", 20), ("right", 72)):
        _sub(margins, "w:" + side, w=w, type="dxa")
    _sub(tbl_pr, "w:tblLook", val="0000")

    for i, cell in enumerate(table.rows[0].cells):
        cell.width = Twips(widths[i])
        cell.text = ""
        add_runs(cell.paragraphs[0], header[i], TABLE_PT)
        cell.paragraphs[0].style = doc.styles["Table Head"]

    for brow in body:
        cells = table.add_row().cells
        for i in range(ncol):
            cells[i].width = Twips(widths[i])
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.style = doc.styles["Table text"]
            add_runs(p, brow[i] if i < len(brow) else "", TABLE_PT)

    for row in table.rows:                            # cantSplit precedes tblHeader in trPr
        _sub(row._tr.get_or_add_trPr(), "w:cantSplit")
    _sub(table.rows[0]._tr.get_or_add_trPr(), "w:tblHeader")   # repeat header across pages
    para(doc, "Normal")


# --------------------------------------------------------------------------- parsing

def read_blocks(lines: list[str], start: int, stop: int) -> list[str]:
    """Group `lines[start:stop]` into paragraphs on blank lines, unwrapping each."""
    out, buf = [], []
    for raw in lines[start:stop]:
        if raw.strip():
            buf.append(raw.strip())
        elif buf:
            out.append(" ".join(buf))
            buf = []
    if buf:
        out.append(" ".join(buf))
    return out


def parse_front_matter(lines: list[str]) -> dict:
    """Pull title, authors, abstract and keywords out of the manuscript's head.

    Authors are matched to affiliations by the superscript marker each carries, which is
    the only link the Markdown records between them.
    """
    title = next(ln[2:].strip() for ln in lines if ln.startswith("# "))

    names_at = next(i for i, ln in enumerate(lines)
                    if ln.startswith("**") and ln.rstrip().endswith("**"))
    names = []
    for chunk in re.sub(r"\*+", "", lines[names_at]).split(","):
        chunk = chunk.replace("\\", "").strip()
        marker = next((c for c in chunk if c in SUPERS), None)
        cleaned = "".join(c for c in chunk if c not in SUPERS + "*").strip()
        if cleaned:
            names.append((marker, cleaned))

    affil: dict[str, tuple[str, str]] = {}
    for ln in lines[names_at:]:
        if ln[:1] and ln[:1] in SUPERS:
            text = ln[1:].strip()
            email = ""
            m = re.search(r"Email:\s*(\S+)", text)
            if m:
                email = m.group(1).rstrip(".")
                text = text[: m.start()].strip().rstrip(".")
            affil[ln[0]] = (text, email)

    corresponding = next(
        (re.sub(r"[*\\]", "", ln).strip() for ln in lines[names_at:]
         if "Corresponding author" in ln), "")

    abs_at = next(i for i, ln in enumerate(lines) if ln.strip() == "## Abstract")
    kw_at = next(i for i, ln in enumerate(lines) if ln.startswith("**Keywords:**"))
    body_at = next(i for i in range(kw_at, len(lines)) if lines[i].startswith("## "))

    keywords = " ".join(read_blocks(lines, kw_at, kw_at + 4)[:1])
    keywords = re.sub(r"^\*\*Keywords:\*\*\s*", "", keywords)

    return {
        "title": title,
        "authors": [{"name": n, "affiliation": affil.get(m, ("", ""))[0],
                     "email": affil.get(m, ("", ""))[1]} for m, n in names],
        "corresponding": corresponding,
        "abstract": read_blocks(lines, abs_at + 1, kw_at),
        "keywords": keywords,
        "body_at": body_at,
    }


# --------------------------------------------------------------------------- title block

def build_title_block(doc, front: dict) -> None:
    """Rewrite the template's title table for this paper's title and author list."""
    tbl = doc.element.body.find(qn("w:tbl"))
    grid = tbl.find(qn("w:tblGrid"))
    rows = tbl.findall(qn("w:tr"))
    n = max(len(front["authors"]), 1)
    share = PAGE_TWIPS // n

    for col in grid.findall(qn("w:gridCol")):
        grid.remove(col)
    for _ in range(n):
        _sub(grid, "w:gridCol", w=share)

    # Row 1 -- the title, spanning every author column.
    title_tc = rows[0].find(qn("w:tc"))
    title_pr = title_tc.find(qn("w:tcPr"))
    title_pr.find(qn("w:tcW")).set(qn("w:w"), str(PAGE_TWIPS))
    span = title_pr.find(qn("w:gridSpan"))
    (span if span is not None else _sub(title_pr, "w:gridSpan")).set(qn("w:val"), str(n))
    title_ps = title_tc.findall(qn("w:p"))
    for p in title_ps[1:]:
        title_tc.remove(p)
    tp = Paragraph(title_ps[0], doc)
    for run in list(tp.runs):
        run._r.getparent().remove(run._r)
    add_runs(tp, front["title"])

    # Row 2 -- one cell per author: name, affiliation in italic, e-mail.
    cells = rows[1].findall(qn("w:tc"))
    while len(cells) > n:
        rows[1].remove(cells.pop())
    while len(cells) < n:
        clone = copy.deepcopy(cells[-1])
        rows[1].append(clone)
        cells.append(clone)
    for tc, author in zip(cells, front["authors"]):
        tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(share))
        for p in tc.findall(qn("w:p"))[1:]:
            tc.remove(p)
        ap = Paragraph(tc.find(qn("w:p")), doc)
        for run in list(ap.runs):
            run._r.getparent().remove(run._r)
        ap.add_run(author["name"])
        ap.add_run().add_break()
        ap.add_run(author["affiliation"]).italic = True
        if author["email"]:
            ap.add_run().add_break()
            ap.add_run(author["email"])
        for run in ap.runs:
            run.font.size = Pt(10.5)


# --------------------------------------------------------------------------- rendering

def clear_boilerplate(doc) -> None:
    """Drop the template's instructional text, keeping its title block and sectPrs."""
    body = doc.element.body
    children = list(body)
    cut = next(i for i, ch in enumerate(children)
               if ch.tag == qn("w:p")
               and ch.find(qn("w:pPr")) is not None
               and ch.find(qn("w:pPr")).find(qn("w:sectPr")) is not None)
    for child in children[cut + 1:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def render(md_path: Path, out_path: Path, template: Path) -> list[str]:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    front = parse_front_matter(lines)

    doc = Document(str(template))
    clear_boilerplate(doc)
    build_title_block(doc, front)

    if front["corresponding"]:
        # Goes in the single-column title section, under the authors -- the two-column
        # body starts at the abstract.
        p = Paragraph(doc.element.body.findall(qn("w:p"))[0], doc)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, front["corresponding"], size=9)
        for run in p.runs:
            run.italic = True

    for k, block in enumerate(front["abstract"]):
        p = para(doc, "Abstract", ("Abstract — " if k == 0 else "") + block)
        p.paragraph_format.space_after = Pt(3)   # structured abstract: keep the labels apart
    para(doc, "Abstract", "")
    para(doc, "Abstract", "Index Terms — " + front["keywords"])

    i, in_refs, placeholders = front["body_at"], False, []
    in_list, num_id = False, None
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        # A numbered item may only *begin* a block or continue a list. Without this a
        # wrapped prose line -- "...pushing risks too far toward 0 and / 1. The CatBoost
        # arms sit at 0.923..." -- is silently turned into list item 1.
        opens_block = i == 0 or not lines[i - 1].strip()

        if not stripped or re.fullmatch(r"-{3,}|\*{3,}", stripped):
            i += 1
            continue
        was_list, in_list = in_list, False

        # ---- headings: the Markdown's own numbering is preserved (see module docstring)
        m = re.match(r"^(#{2,6})\s+(.*)$", stripped)
        if m:
            lvl, text = len(m.group(1)), m.group(2)
            if lvl == 2:
                in_refs = text.strip().lower().endswith("references")
                para(doc, "Heading 1", text)
            elif lvl == 3:
                para(doc, "Heading 2", text)
            else:
                p = para(doc, "Body Text Indent", text, first_line=0.0)
                p.paragraph_format.space_before = Pt(6)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_with_next = True
                for run in p.runs:
                    run.bold = True
            i += 1
            continue

        # ---- pipe table written out in the manuscript itself
        if (stripped.startswith("|") and i + 1 < len(lines)
                and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1])):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            add_table(doc, block, COL_TWIPS)
            continue

        # ---- blockquote: a table/figure placeholder, or a pull-quote
        if stripped.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip(">").strip())
                i += 1
            joined = " ".join(block)
            marker = re.search(r"\[(TABLE|FIGURE) (\d+) HERE\]", joined)
            if marker:
                placeholders.append(f"{marker.group(1).title()} {marker.group(2)}")
                add_placeholder(doc, joined)
            else:
                para(doc, "BlockQuote", joined)
            continue

        # ---- list item, possibly wrapped over continuation lines
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", raw)
        if m and (opens_block or was_list):
            indent, marker, text = m.group(1), m.group(2), m.group(3)
            i += 1
            while (i < len(lines) and lines[i].strip()
                   and not re.match(r"^\s*([-*]|\d+\.)\s+", lines[i])
                   and not lines[i].strip().startswith(("#", "|", ">"))
                   and lines[i].startswith((" ", "\t"))):
                text += " " + lines[i].strip()
                i += 1
            numbered = marker[0].isdigit()
            p = para(doc, "NumberedList" if numbered else "BulletList", text)
            if len(indent) >= 2:
                p.paragraph_format.left_indent = Inches(0.55)
            if numbered:
                if was_list and num_id is not None:
                    apply_numbering(p, num_id)
                else:
                    num_id = restart_numbering(doc, p)
            in_list = True
            continue

        # ---- paragraph: join wrapped lines until a blank line or a block starter
        buf = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            # Bullets may interrupt a paragraph; a numbered line may not, for the reason
            # given above -- every list in these manuscripts is blank-line separated.
            if (not nxt or nxt.startswith(("#", "|", ">"))
                    or re.fullmatch(r"-{3,}|\*{3,}", nxt)
                    or re.match(r"^\s*[-*]\s+", lines[i])):
                break
            buf.append(nxt)
            i += 1
        text = " ".join(buf)
        if in_refs and re.match(r"^\[\d+\]", text):
            p = para(doc, "Heading 3", text)         # template alias: 'References Text'
            p.paragraph_format.space_after = Pt(6)
        else:
            para(doc, "Body Text Indent", text)

    doc.save(str(out_path))
    return placeholders


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("markdown", nargs="?", default="CAFE_Hybrid_Manuscript.md")
    ap.add_argument("-o", "--out", default=None)
    ap.add_argument("--template", default=str(TEMPLATE))
    args = ap.parse_args()

    md = Path(args.markdown)
    if not md.is_absolute():
        md = REPO / md
    if not md.exists():
        print(f"not found: {md}")
        return 1

    out = (Path(args.out) if args.out else
           md.with_name(md.stem.split("_Manuscript")[0] + "_TemplateFormat.docx"))
    if not out.is_absolute():
        out = REPO / out

    template = Path(args.template)
    if not template.is_absolute():
        template = REPO / template

    placeholders = render(md, out, template)
    print(f"wrote {out.relative_to(REPO)}  ({out.stat().st_size / 1024:.0f} KB)")
    print(f"  template     : {template.name}")
    print(f"  placeholders : {len(placeholders)} left for manual insertion "
          f"({', '.join(placeholders)})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
